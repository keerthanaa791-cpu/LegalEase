import html
import io
import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "assets" / "logo.png"
FONT_NAME = "Times New Roman"

def sanitize_text(text: str) -> str:
    replacements = {"“": '"', "”": '"', "‘": "'", "’": "'", "—": "-", "–": "-", "•": "-"}
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()

def split_sections(text: str) -> list[tuple[str, list[str]]]:
    lines = [line.strip() for line in sanitize_text(text).splitlines()]
    sections: list[tuple[str, list[str]]] = []
    heading = ""
    body: list[str] = []
    for line in lines:
        if not line:
            continue
        is_heading = bool(re.match(r"^(\d+\.?\s+|[A-Z][A-Z\s/&-]{4,}$)", line)) and len(line) < 140
        if is_heading:
            if heading or body:
                sections.append((heading, body))
            heading, body = line, []
        else:
            body.append(line)
    if heading or body:
        sections.append((heading, body))
    return sections or [("Document", [sanitize_text(text)])]

def format_docx(text: str, doc_type: str, branding_name: str = "LegalEase", terms: Optional[list[str]] = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME; normal.font.size = Pt(11)

    if LOGO.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.25))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(doc_type.upper()); r.bold = True; r.font.name = FONT_NAME; r.font.size = Pt(16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(branding_name).italic = True

    for heading, body in split_sections(text):
        if heading:
            p = doc.add_paragraph(); r = p.add_run(heading); r.bold = True; r.font.size = Pt(12)
        for line in body:
            if line.startswith("-"):
                p = doc.add_paragraph(style="List Bullet"); p.add_run(line.lstrip("- "))
            else:
                doc.add_paragraph(line)

    if terms:
        doc.add_paragraph().add_run("Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"; table.rows[0].cells[1].text = "Term"
        for idx, term in enumerate(terms, 1):
            cells = table.add_row().cells; cells[0].text = str(idx); cells[1].text = term

    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{branding_name} - Draft generated with AI - Review before use")
    out = io.BytesIO(); doc.save(out); return out.getvalue()

class LegalEasePDF(FPDF):
    def __init__(self, branding_name: str):
        super().__init__(); self.branding_name = branding_name
    def header(self):
        if LOGO.exists():
            try: self.image(str(LOGO), x=90, y=8, w=30)
            except Exception: pass
        self.set_y(22)
        self.set_font("Times", "B", 9); self.cell(0, 8, self.branding_name, align="C")
        self.ln(7)
    def footer(self):
        self.set_y(-15); self.set_font("Times", "I", 8)
        self.cell(0, 10, f"{self.branding_name} - AI draft - Page {self.page_no()}", align="C")

def format_pdf(text: str, doc_type: str, branding_name: str = "LegalEase") -> bytes:
    pdf = LegalEasePDF(branding_name); pdf.set_auto_page_break(True, margin=18); pdf.add_page()
    pdf.set_font("Times", "B", 16); pdf.multi_cell(0, 9, doc_type.upper(), align="C"); pdf.ln(4)
    for heading, body in split_sections(text):
        if heading:
            pdf.set_font("Times", "B", 12); pdf.multi_cell(0, 7, heading)
        pdf.set_font("Times", "", 11)
        for line in body:
            if line.startswith("-"):
                pdf.multi_cell(0, 6, "- " + line.lstrip("- "))
            else:
                pdf.multi_cell(0, 6, line)
            pdf.ln(1)
    return bytes(pdf.output(dest="S"))

def format_html_preview(text: str) -> str:
    safe = html.escape(sanitize_text(text))
    safe = re.sub(r"(^|\n)(\d+\.?\s+[^\n]+)", r"\1<h3>\2</h3>", safe)
    safe = safe.replace("\n", "<br>")
    return f'<div class="legal-preview">{safe}</div>'
